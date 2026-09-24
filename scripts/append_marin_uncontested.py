# -*- coding: utf-8 -*-
import openpyxl

XLSX_PATH = "/Users/ciarazavala/Desktop/Ballot Preview Backfill/Measure_Candidate Backfill Info.xlsx"
COUNTY = "Marin County"
URL = "https://www.marincounty.gov/departments/elections/november-3-2026-general-election/information-and-about-candidates-110326/local-candidates-not-ballot-110326"

BASE_FLAG = "Uncontested - registrar lists this on the \"Local candidates NOT ON ballot\" page (filed candidates do not exceed available seats, so no election is held)"
CROSS_COUNTY_FLAG = BASE_FLAG + ' ; Cross-county jurisdiction: Sonoma County Junior College District, not Marin County - verify before publishing under "Marin County"'

RACES = [
    ("Marin Community College District Governing Board Member, Trustee Area 1", [("Crow, Suzanne Brown", "Incumbent")]),
    ("Marin Community College District Governing Board Member, Trustee Area 3", [("Treanor, Wanden Patricia", "Incumbent")]),
    ("Marin Community College District Governing Board Member, Trustee Area 6", [("Conti, Diana", "Incumbent")]),
    ("Marin Community College District Governing Board Member, Trustee Area 7 - Short Term", [("Deakyne, Jessica", "Chief Operating Officer")]),
    ("Sonoma County Junior College District Governing Board Member, Trustee Area 2", [("Fishman, Margaret", "Incumbent")]),
    ("San Rafael City Schools Board of Education Governing Board Member, Trustee Area 5, Short Term", [("Rodas-Dias, Kennedy R.", "Appointed Incumbent")]),
    ("San Rafael City Schools Board of Education Governing Board Member, Trustee Area 2", [("Lau, Wing Keung Jason", "Appointed Trustee / Educator")]),
    ("San Rafael City Schools Board of Education Governing Board Member, Trustee Area 4", [("Martin, Carolina", "Incumbent")]),
    ("Bolinas-Stinson Union School District Governing Board Member", [("Dar, Arianne Z.", "School Board Trustee"), ("Tonski, Jacob", "Incumbent")]),
    ("Kentfield School District Governing Board Member", [("Goldman, Brent J.", "Education Consultant / Parent"), ("Portnoy, Zachary", "Education Technology Advisor"), ("Wilda, Traci", "P T A President / Businesswoman")]),
    ("Laguna Joint School District Governing Board Member", [("Hess, Sharon Beretta", "Incumbent"), ("Hess, Daniel", "Retired")]),
    ("Larkspur-Corte Madera School District Governing Board Member", [("Dinday, Jessica", "Appointed Incumbent"), ("Movafaghi, Amir", "Incumbent"), ("Woodring, Ryan", "Clinical Social Worker")]),
    ("Mill Valley School District Governing Board Member", [("Katz, Natalie", "Incumbent"), ("Nakatani, Sharon", "Incumbent"), ("Yoo, Yunhee", "Incumbent")]),
    ("Nicasio School District Governing Board Member", [("Burton, Mark", "Incumbent")]),
    ("Novato Unified School District Governing Board Member, Trustee Area 5", [("Gasson, Diane", "Incumbent")]),
    ("Novato Unified School District Governing Board Member, Trustee Area 6", [("Rogoff, Theresa Odisio", "Realtor / Parent / Educator")]),
    ("Ross Valley School District Governing Board Member", [("Landles-Cobb, Christina", "Incumbent"), ("Marsh, Anna", "Appointed Incumbent")]),
    ("Sausalito Marin City School District Governing Board Member - Short Term", [("Graham, Jamal", "Appointed Incumbent")]),
    ("Shoreline Unified School District Governing Board Member, Trustee Area 1", [("Faure, Buddy", ""), ("Hernandez, Yezenia A.", "Parent")]),
    ("Shoreline Unified School District Governing Board Member, Trustee Area 2", [("Healy, Jane", "Incumbent")]),
    ("City of Novato Councilmember District 5 - Short Term", [("Karkal, Sandeep", "Appointed City Councilmember")]),
    ("Marinwood Community Services District Director", [("Oyserman, Sivan", "Incumbent"), ("Shea, William", "Director, Marinwood Community Services District")]),
    ("Muir Beach Community Services District Director", [("Hills, Leighton", "Incumbent"), ("Murray, Christine", "Incumbent"), ("Shaffer, Steven", "Incumbent")]),
    ("Tamalpais Community Services District Director", [("Brown, Jeffrey", "Incumbent"), ("Keon, Pam", "Community Volunteer")]),
    ("Tomales Village Community Services District Director", [("Bonini, William A.", ""), ("O'Neill, Dru Fallon", "Incumbent"), ("Ward, John", "Incumbent")]),
    ("Bolinas Fire Protection District Director", [("Molesworth, Claire", "Incumbent"), ("Torrey, Nancy", "Incumbent")]),
    ("Kentfield Fire Protection District Director", [("Naso, Ron", "Incumbent"), ("Evergettis, Barry", "Incumbent"), ("Ryan, Dennis", "Appointed Incumbent")]),
    ("Novato Fire Protection District Director, District 4", [("Goines, Bruce F", "Incumbent")]),
    ("Novato Fire Protection District Director, District 5", [("Davis, Bill", "Incumbent")]),
    ("Sleepy Hollow Fire Protection District Director", [("Shortall, Richard", "Incumbent")]),
    ("Sleepy Hollow Fire Protection District Director - Short Term", [("Gauna, Jennifer", "Appointed Incumbent")]),
    ("Stinson Beach Fire Protection District Director", [("Poler, Ariel", "")]),
    ("Stinson Beach Fire Protection District Director - Short Term", [("Bertrand, Christophe", "Appointed Incumbent")]),
    ("Tiburon Fire Protection District Director", [("Ho, Joy", "Director, Tiburon Fire Protection District"), ("Jones, Richard", "Director, Tiburon Fire Protection District"), ("Woodford, Cheryl", "Director, Tiburon Fire Protection District")]),
    ("Marin Healthcare District Director, Division 4", [("Rienks, Jennifer", "Incumbent")]),
    ("Inverness Public Utility District Director", [("Press, David", "Incumbent"), ("Longstreth, John", "Retired")]),
    ("Mesa Park Recreation District Director", [("Dunne, Simon", ""), ("Lowrance, Ben", "")]),
    ("Mesa Park Recreation District Director - Short Term", [("Suda, Dan", "")]),
    ("Marin Resource Conservation District Director", [("Meral, Jerry", "Incumbent"), ("Williams, Melissa", "Appointed Incumbent")]),
    ("Homestead Valley Sanitary District Director", [("Leibof, Allan", "Incumbent"), ("Saltzman, Alan", "Incumbent")]),
    ("Homestead Valley Sanitary District Director - Short Term", [("Alexandris, Penelope", "Appointed Incumbent")]),
    ("Las Gallinas Valley Sanitary District Director", [("Clark, Megan", "Incumbent"), ("Lavrov, Nicholas", "Director, Las Gallinas Valley Sanitary District"), ("Yezman, Crystal", "Incumbent")]),
    ("Novato Sanitary District Director, Division 2", [("Bentley, Dennis E.", "Incumbent")]),
    ("Novato Sanitary District Director, Division 3", [("Fuette, Tim", "Incumbent")]),
    ("Novato Sanitary District Director, Division 5", [("Jacobs, Sherri", "Retired Business Manager")]),
    ("Novato Sanitary District Director, Division 1 - Short Term", [("Tarantino, Jeffrey J.", "Appointed Incumbent")]),
    ("Sausalito-Marin City Sanitary District Director", [("Rheiner, Dan J.", "Incumbent"), ("Rycerski, Barbara", ""), ("Thornton, Shirley", "Incumbent")]),
    ("Sausalito-Marin City Sanitary District Director - Short Term", [("McKibben, James", "Appointed Incumbent")]),
    ("Sanitary District No. 5 Director", [("Benediktsson, Catharine", "Incumbent"), ("Moody, Tod", "Incumbent")]),
    ("Marin Municipal Water District Director Division 1", [("Samson, Matthew", "Marin Water Director")]),
    ("Marin Municipal Water District Director Division 3", [("Khush, Ranjiv", "Director, Marin Municipal Water District")]),
    ("Marin Municipal Water District Director Division 4", [("Lovejoy, Erika", "Environmental Program Director")]),
    ("North Marin Water District Director, Division 1", [("Eichstaedt, Ken", "Incumbent")]),
    ("Stinson Beach County Water District Director", [("Baskin, Lawrence 'Larry'", ""), ("Zell, Jim", "Director, Stinson Beach County Water District")]),
]

wb = openpyxl.load_workbook(XLSX_PATH)
ws_c = wb['Candidates']

c_row = 2
while ws_c.cell(row=c_row, column=1).value:
    c_row += 1

n = 0
for race, cands in RACES:
    flag = CROSS_COUNTY_FLAG if race.startswith("Sonoma County Junior College District") else BASE_FLAG
    for name, occ in cands:
        ws_c.cell(row=c_row, column=1, value=COUNTY)
        ws_c.cell(row=c_row, column=2, value=race)
        ws_c.cell(row=c_row, column=3, value=name)
        ws_c.cell(row=c_row, column=4, value=occ)
        ws_c.cell(row=c_row, column=5, value=flag)
        ws_c.cell(row=c_row, column=6, value=URL)
        c_row += 1
        n += 1

wb.save(XLSX_PATH)
print(f"Marin uncontested rows appended: {n} across {len(RACES)} races")
