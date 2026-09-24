# -*- coding: utf-8 -*-
import openpyxl

XLSX_PATH = "/Users/ciarazavala/Desktop/Ballot Preview Backfill/Measure_Candidate Backfill Info.xlsx"
COUNTY = "San Joaquin County"
MAIN_URL = "https://www.sjgov.org/department/rov/election-information/current_election"

# (letter, jurisdiction, description, source_url, flag)
MEASURES = [
    ("G", "Galt Joint Union High School District",
     "Galt Joint Union High School District, Repair/Safety and Vocational Training Measure. To upgrade classrooms to maintain quality education/school safety by preparing students for skilled trades including welding/construction/plumbing; repairing leaky roofs; removing asbestos/mold; shall Galt Joint Union High School District’s measure be adopted, authorizing $43,000,000 in bonds at legal rates, levying $30 per $100,000 of assessed valuation, generating $3,100,000 annually, while bonds are outstanding, prohibiting funds for administrators salaries/pensions, requiring independent citizens’ oversight, audits, all funds locally controlled?",
     MAIN_URL + " ; https://elections.saccounty.gov/us/en/election-information/current-measure-information.html#gsc.tab=0",
     'Cross-county jurisdiction: Galt Joint Union High School District measure is led by Sacramento County (per San Joaquin County’s own page: "Sacramento County is the lead county") - verify before publishing under "San Joaquin County"'),
    ("T", "Lodi Unified School District",
     "WITH NO ESTIMATED INCREASE IN CURRENT TAX RATES, shall Lodi Unified School District’s measure to repair leaky roofs, plumbing, deteriorating electrical, gas lines, heating, and air conditioning; improve school security; and upgrade career training and college preparation classrooms be adopted, authorizing $195 million of bonds with legal rates, average annual levies below $22 per $100,000 of assessed valuation ($11.4 million annually) while outstanding, annual audits, independent citizen oversight, full disclosure of all spending, and all funds locally controlled?",
     MAIN_URL + " ; https://www.sjgov.org/docs/default-source/registrar-of-voters-documents/elections/gubernatorial-general-election-2026/candidate-services/notice-of-a-measure-t_eng_final.pdf",
     ""),
    ("U", "Escalon Unified School District",
     "To provide safe, modern schools by upgrading/repairing aging classrooms, science labs, career technology and school facilities to support college/career readiness in math, science, technology, arts, engineering, and skilled trades; and fixing deteriorating roofs, plumbing, HVAC, and electrical systems; shall Escalon Unified School District’s measure authorizing $26,000,000 in bonds at legal rates be adopted, levying approximately $18 per $100,000 assessed value ($1,700,000 annually) while bonds are outstanding, with independent audits, citizen oversight and all money locally controlled?",
     MAIN_URL + " ; https://www.sjgov.org/docs/default-source/registrar-of-voters-documents/elections/gubernatorial-general-election-2026/candidate-services/notice-of-a-measure-u_eng_final.pdf",
     ""),
    ("V", "Lincoln Unified School District",
     "Lincoln Unified School District Facilities and Safety Measure. To improve the quality of education; make safety/security improvements; provide classrooms and labs for job training and career readiness; replace roofs and HVAC systems; restrooms and facilities improvements; shall Lincoln Unified School District’s measure authorizing $79,000,000 of bonds be adopted at legal rates, estimated levies averaging $39.07 per $100,000 assessed value ($4,471,000 annually) while bonds are outstanding, with citizens’ oversight, audits, no money for salaries and funding only for LUSD schools?",
     MAIN_URL + " ; https://www.sjgov.org/docs/default-source/registrar-of-voters-documents/elections/gubernatorial-general-election-2026/candidate-services/v_english-rn-44458773.pdf",
     ""),
]

wb = openpyxl.load_workbook(XLSX_PATH)
ws_m = wb['Measures']

m_row = 2
while ws_m.cell(row=m_row, column=1).value:
    m_row += 1

n = 0
for letter, jurisdiction, desc, url, flag in MEASURES:
    ws_m.cell(row=m_row, column=1, value=COUNTY)
    ws_m.cell(row=m_row, column=2, value=f"Measure {letter}")
    ws_m.cell(row=m_row, column=3, value=jurisdiction)
    ws_m.cell(row=m_row, column=4, value=desc)
    ws_m.cell(row=m_row, column=5, value=flag)
    ws_m.cell(row=m_row, column=6, value=url)
    m_row += 1
    n += 1

wb.save(XLSX_PATH)
print(f"Measures rows written: {n}")
