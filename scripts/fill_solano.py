# -*- coding: utf-8 -*-
import openpyxl

XLSX_PATH = "/Users/ciarazavala/Desktop/Ballot Preview Backfill/Measure_Candidate Backfill Info.xlsx"
CANDIDATES_SOURCE = "https://content.solanocounty.gov/sites/default/files/2026-07/CANDIDATE_FILED_LOG_Nov_3_2026_7.13.26.pdf"
MEASURES_SOURCE = "https://www.solanocounty.gov/government/registrar-voters/current-election-information"
COUNTY = "Solano County"

MEASURES = [
    ("Measure B", "Winters Joint Unified School District",
     "To improve the quality of education; repair or replace leaky roofs; modernize/renovate outdated classrooms, restrooms and school facilities; and replace deteriorating plumbing and sewer systems; shall Winters Joint Unified School District's measure authorizing $29,900,000 in bonds, at legal interest rates, and levying approximately $49 per $100,000 of assessed valuation (raising $1,600,000 annually) while bonds are outstanding, be adopted, with a Board-appointed citizens' oversight committee and annual independent audits to assure proper expenditure of funds?",
     "Cross-county jurisdiction (Yolo and Solano counties) - please verify"),
    ("Measure E", "County of Solano (unincorporated area)",
     "Measure __: Solano County Unincorporated Business License Tax. Shall Solano County update its business license tax to raise the commercial wind turbine rate ($0.00003 to $0.00008 per kWh) and extend the tax to commercial solar energy systems ($0.00008 per kWh), battery energy storage systems ($1.50 per MWh), natural gas extraction facilities ($0.30 per McF), and data centers ($5.00 per building square foot), with annual inflation adjustments, generating approximately an additional $353,000 annually, until ended by the voters?",
     ""),
    ("Measure H", "County of Solano (unincorporated area)",
     "Measure H: Solano County Unincorporated Transient Occupancy Tax. To fund local programs and essential public services, shall Solano County increase its existing Transient Occupancy Tax, paid by tourists and others staying overnight at short-term lodging facilities in unincorporated areas (excluding cities), from 5% to 12% of the total rent charged, providing approximately an additional $58,500 annually, until ended by the voters?",
     ""),
    ("Measure O", "City of Rio Vista",
     "Measure \"O\" - Extension of General City Services Tax Measure. Shall the existing three-quarter (3/4%) percent transactions and use tax measure to support the City of Rio Vista's general services, including police, fire, parks, and other unrestricted general fund expenditures, be extended for three years, commencing on April 1, 2027, and which is projected to provide between $1,398,000 to $1,467,000 of revenue annually for the City?",
     ""),
    ("Measure P", "City of Fairfield",
     "City of Fairfield Measure P Update Measure. To provide and protect funding for Fairfield's services, such as police, fire, 911 emergency response, neighborhood safety, street and sidewalk repair, homelessness prevention and response, and general governmental use, shall an ordinance be adopted to extend the City's sales tax and increase the rate by up to one-cent per dollar until ended by voters, adding approximately $24 million annually, with independent citizen oversight, annual audits, public spending reports, and all funds locally controlled?",
     ""),
    ("Measure Q", "City of Fairfield",
     "City of Fairfield Appointive City Clerk Measure. Shall the Office of City Clerk be appointive?",
     ""),
    ("Measure U", "Dixon Unified School District",
     "To fix deteriorating roofs, plumbing, gas, and electrical systems, remove asbestos/lead pipes, repair/upgrade aging classrooms, science labs, technology and school facilities to support college/career readiness in math, science, technology, arts, engineering, and skilled trades; shall Dixon Unified School District's measure authorizing $49,000,000 in bonds at legal rates be adopted, levying approximately $30 per $100,000 assessed value ($3,000,000 annually) while bonds are outstanding, with citizen oversight, independent audits, and all money locally controlled?",
     ""),
    ("Measure V", "City of Vacaville",
     "Vacaville Essential Services Measure of 2026. Shall the citizens of Vacaville adopt a measure maintaining the City's financial stability and essential services—including 911 emergency response; fire and ambulance protection; crime prevention; parks and recreation; safe public spaces; supporting local businesses; repairing streets; and other general City services—by establishing an additional one-cent sales tax providing approximately twenty-eight million dollars annually, with local control, citizens oversight, annual audits, and public reporting, until ended by voters?",
     ""),
    ("Measure X", "City of Benicia",
     "Measure X: Benicia Business License Modernization. Shall the City of Benicia adopt a Business License Tax measure, exempting businesses with $100,000 or less in annual receipts, setting tax rates by business category ranging from $0.25 to $1.25 per $1,000 in receipts, capping those annual taxes at $1,500 in the first year with 10% annual increases until the cap is $25,000, with a separate Marine Terminal rate of $3 per $100 in receipts, generating approximately $1,612,000 in the first year, until repealed by voters?",
     ""),
    ("Measure Y", "City of Benicia",
     "Measure Y: Limited City Charter. Shall the measure, adopting a limited City Charter for the City of Benicia that allows the real property transfer tax in Measure Z to take effect if voters approve both Measures Y and Z, and continues the City's existing council-manager form of government, be adopted?",
     "Linked to Measure Z - Measure Z's transfer tax only takes effect if both Y and Z are approved"),
    ("Measure Z", "City of Benicia",
     "Measure Z: Real Property Transfer Tax. Shall the Measure, to primarily ensure newly constructed property sales contribute to maintaining city services and infrastructure, establish a real property transfer tax that will not apply to existing homes that sell for $2 million or less, but apply to other property sales at rates of 0.4% for sales $2,000,000 or less, 0.6% for sales $2,000,001 - $10,000,000, and 0.8% for sales over $10,000,000, providing approximately $230,000 annually, until ended by voters, be adopted?",
     "Linked to Measure Y - only takes effect if both Y and Z are approved"),
]

# (race, seats, [(name, occupation), ...], flag_override_or_None)
RACES = [
    ("Solano County Board of Supervisors - District 4", 1, [
        ("Michael \"Mike\" Silva", "Professor / Vacaville Councilmember"),
        ("John Carli", "Mayor, City of Vacaville"),
    ], None),

    ("Solano County Board of Education - Trustee Area 3", 1, [
        ("Dana Dean", "Incumbent"),
    ], None),
    ("Solano County Board of Education - Trustee Area 4", 1, [
        ("Teresa Lavell", "Solano County Board of Education Trustee Area 4"),
    ], None),
    ("Solano County Board of Education - Trustee Area 6", 1, [
        ("Tamika Hamilton", "Retired USAF Sergeant"),
        ("Kathryn Kedarisetty", "Elementary School Teacher"),
    ], None),
    ("Yolo County Board of Education - Trustee Area 2", 1, [
        ("Melissa Moreno", ""),
    ], "Cross-county jurisdiction - filed in Yolo County per Solano registrar log; ballot designation/challenger data incomplete in this source, please verify directly with Yolo County registrar"),
    ("Yolo County Board of Education - Trustee Area 5", 1, [
        ("Julian Munoz", ""),
    ], "Cross-county jurisdiction - filed in Yolo County per Solano registrar log; ballot designation/challenger data incomplete in this source, please verify directly with Yolo County registrar"),

    ("San Joaquin Delta Community College - Trustee Area 4", 1, [
        ("Charles R. Jennings", "Delta College Trustee"),
        ("Tony Shah", "Retired School Administrator"),
    ], "Cross-county jurisdiction (San Joaquin and Solano counties) - filed in San Joaquin County per Solano registrar log; please verify"),
    ("Solano Community College - Trustee Area 1", 1, [
        ("Karimah Karah", "Financial Advisor / Parent"),
    ], None),
    ("Solano Community College - Trustee Area 2", 1, [
        ("Annie M. Young", "Retired Mathematics Professor"),
    ], None),
    ("Solano Community College - Trustee Area 4", 1, [
        ("Denis Honeychurch", "College Trustee / Attorney"),
        ("Tony Yang", "College Analyst / Administrator"),
    ], None),
    ("Solano Community College - Trustee Area 6", 1, [
        ("Amber M. Cargo-Reed", "Solano Community College Board Trustee Area 6"),
    ], None),

    ("Benicia Unified School District - Trustee Area 1", 1, [
        ("Dan Smith", "Substitute Teacher"),
        ("Barbara Jones", "Retired Teacher"),
    ], None),
    ("Benicia Unified School District - Trustee Area 4", 1, [
        ("Erin Fernandez", "Hospital Billing Manager"),
    ], None),
    ("Benicia Unified School District - Trustee Area 5", 1, [
        ("Amy Hirsh", "Incumbent"),
    ], None),

    ("Dixon Unified School District (at-large)", 3, [
        ("Jewel Fink", "Incumbent"),
        ("Julian Y. Cuevas", "Incumbent"),
        ("John Gabby", "Parent"),
        ("Phil Lockwood", "Chaplain / Wellness Coach"),
        ("Michael Ceremello Jr.", "Investigative Journalist"),
        ("Cheryl Sommers", "Attorney"),
        ("Michael Monson", "Electrician / Parent"),
    ], None),

    ("Fairfield-Suisun Unified School District - Trustee Area 4", 1, [
        ("Tom Ewen", "Carpenter / Union Trustee"),
    ], None),
    ("Fairfield-Suisun Unified School District - Trustee Area 5", 1, [
        ("MaryDelores Wilson", "Mental Health Clinician"),
    ], None),
    ("Fairfield-Suisun Unified School District - Trustee Area 7", 1, [
        ("Kai Eusebio", "Fairfield-Suisun Unified School District, Board Member Trustee Area 7"),
    ], None),

    ("River Delta Unified School District - Trustee Area 1", 1, [
        ("Dan Mahoney", "Incumbent"),
        ("Michael A Mimiaga", "School Transportation Director"),
    ], None),
    ("River Delta Unified School District - Trustee Area 3", 1, [], None),

    ("Travis Unified School District - Trustee Area 1", 1, [], None),
    ("Travis Unified School District - Trustee Area 2", 2, [
        ("Matthew Bidou", "Incumbent"),
        ("Tierra Ekstrom", "Parent"),
    ], None),

    ("Vacaville Unified School District - Trustee Area 1", 1, [
        ("Nancy Dunn", "Vacaville Unified School District Trustee"),
    ], None),
    ("Vacaville Unified School District - Trustee Area 3", 1, [
        ("John Jansen", "Incumbent"),
    ], None),
    ("Vacaville Unified School District - Trustee Area 5", 1, [
        ("Santiago Serrato", "Incumbent"),
    ], None),

    ("Vallejo City Unified School District - Trustee Area 1", 1, [
        ("Clarence John Martin Jr", "Business Owner"),
        ("Judith Lerner", "Retired Teacher"),
        ("Reba Jill Brown", "Tax Preparer"),
    ], None),
    ("Vallejo City Unified School District - Trustee Area 3", 1, [
        ("Rizal \"Mr. Riz\" V. Aliga", "Tennis Coach / Educator"),
        ("Irene Reynolds", "Small Business Owner"),
    ], None),
    ("Vallejo City Unified School District - Trustee Area 5", 1, [
        ("John Fox", "Teacher"),
    ], None),

    ("Winters Joint Unified School District - Trustee Area 2", 1, [
        ("Carrie Green", ""),
    ], "Cross-county jurisdiction - filed in Yolo County per Solano registrar log; ballot designation/challenger data incomplete in this source, please verify directly with Yolo County registrar"),
    ("Winters Joint Unified School District - Trustee Area 3", 1, [
        ("Joedy Michael", ""),
    ], "Cross-county jurisdiction - filed in Yolo County per Solano registrar log; ballot designation/challenger data incomplete in this source, please verify directly with Yolo County registrar"),

    ("Cordelia Fire Protection District (at-large)", 3, [
        ("Mark Giugni", "Appointed Incumbent"),
        ("Jim Frische", "Incumbent"),
        ("Dave Carpenter", "Retired Fire Chief"),
    ], None),

    ("Rural North Vacaville Water District (at-large)", 3, [
        ("Eileen Uthe-Smith", "Appointed Incumbent"),
        ("Alan Hanger", "Incumbent"),
        ("Anne Putney", "Government Management Analyst"),
    ], None),

    ("Solano Irrigation District - Division Director 1", 1, [
        ("John D. Kluge", "Incumbent"),
    ], None),
    ("Solano Irrigation District - Division Director 5", 1, [
        ("Derrick Lum", "Incumbent"),
    ], None),

    ("City of Benicia - Councilmember (at-large)", 2, [
        ("Terry Scott", "Incumbent"),
        ("Kari Birdseye", "Incumbent"),
        ("Justin Zollars", "Software Engineer"),
    ], None),

    ("City of Dixon - Councilmember, District 1", 1, [
        ("Jim Ernest", "Incumbent"),
        ("Emily Schroeder", "Legislative Assistant"),
    ], None),
    ("City of Dixon - Councilmember, District 2", 1, [
        ("Thom Bogue", "Councilmember / Business Owner"),
        ("Christina M. Cerna", "Education Compliance Specialist"),
    ], None),
    ("City of Dixon - Elected City Clerk (at-large)", 1, [
        ("Lupe Ruiz", "Administrative City Clerk"),
    ], None),

    ("City of Fairfield - Mayor (at-large)", 1, [
        ("Pam Bertani", "Mother / Councilmember / Attorney"),
        ("Scott Tonnesen", "Small Businessman / Councilmember"),
        ("Nikila Walker Gibson", "Business Owner"),
        ("George Kennedy", "Small Business Owner"),
    ], None),
    ("City of Fairfield - Councilmember, District 1", 1, [
        ("K. Patrice Williams", "Incumbent"),
        ("Nora Dizon", "Real Estate Broker"),
    ], None),
    ("City of Fairfield - Councilmember, District 3", 1, [
        ("Doug Carr", "Incumbent"),
        ("Robert Marin", "Retired Police Officer"),
        ("Susan Bush", "Independent Contractor / Homemaker"),
    ], None),
    ("City of Fairfield - Councilmember, District 5", 1, [
        ("Doriss Panduro", "Incumbent"),
        ("Lynda Davis-Robinson", "Educator / Pastor / Mother"),
    ], None),

    ("City of Rio Vista - Councilmember (at-large)", 2, [
        ("Rick Dolk", "Incumbent"),
        ("Robert Butler", "Retired USMC SGT"),
        ("Gloria J. Thibeaux", "Retired HR Manager"),
        ("Leslie Codling-Dichet", "Human Resources Manager"),
    ], None),

    ("City of Suisun City - Mayor, Full Term (at-large)", 1, [
        ("Alma Hernandez", "Incumbent"),
        ("Chauncey Banks", "Entrepreneur"),
        ("Michael Jefferson", "Political Consultant"),
        ("Lilia Dardon", "Retired Teacher"),
        ("Amit Pal", "Councilmember / Business Owner"),
    ], None),
    ("City of Suisun City - Councilmember (at-large)", 2, [
        ("Princess Washington", "Non-Profit Executive"),
        ("Jenalee Dawson", "Non-Profit Director / Councilmember"),
        ("Xenia Tom", "Home Maker"),
        ("Katrina Garcia", "Small Business Owner"),
    ], None),

    ("City of Vacaville - Mayor (at-large)", 1, [
        ("Sarah Chapman", "Vacaville City Councilmember"),
        ("Roy Stockton", "Councilmember / Sheriff's Sergeant"),
    ], None),
    ("City of Vacaville - Councilmember, District 2", 1, [
        ("Gregory Ritchie II", "Councilmember / Business Owner"),
        ("Tom Chalk", "Retired Assistant Sheriff"),
    ], None),
    ("City of Vacaville - Councilmember, District 4", 1, [
        ("John Vogel", "Public Works Supervisor"),
        ("Davonna Finley", "Community Relations Director"),
    ], None),
    ("City of Vacaville - Councilmember, District 6", 1, [
        ("Jeanette Wylie", "Incumbent"),
        ("Kevin Puett", "Retired Associate Director"),
    ], None),

    ("City of Vallejo - Councilmember, District 2", 1, [
        ("Diosdado \"JR\" Matulac", "Councilmember / Restorative Specialist"),
    ], None),
    ("City of Vallejo - Councilmember, District 4", 1, [
        ("Charles Palmares", "Councilmember / Senior Manager"),
        ("Brianna Rogers", "Community Foundation Strategist"),
        ("Chris Platzer", "Merchant Mariner"),
    ], None),
    ("City of Vallejo - Councilmember, District 5", 1, [
        ("Tara Beasley Stansberry", "Businessowner / HR Recruiter"),
        ("Tanya Hall", "Retired Career Advisor"),
        ("Rebekah Truemper", "Nonprofit Organization Advisor"),
    ], None),
]

wb = openpyxl.load_workbook(XLSX_PATH)
ws_m = wb['Measures']
ws_c = wb['Candidates']

r = 2
while ws_m.cell(row=r, column=1).value:
    r += 1
for name, jurisdiction, desc, flag in MEASURES:
    ws_m.cell(row=r, column=1, value=COUNTY)
    ws_m.cell(row=r, column=2, value=name)
    ws_m.cell(row=r, column=3, value=jurisdiction)
    ws_m.cell(row=r, column=4, value=desc)
    ws_m.cell(row=r, column=5, value=flag)
    ws_m.cell(row=r, column=6, value=MEASURES_SOURCE)
    r += 1
print(f"Wrote {len(MEASURES)} measure rows")

r = 2
while ws_c.cell(row=r, column=1).value:
    r += 1

cand_rows = 0
empty_races = 0
for race, seats, cands, flag_override in RACES:
    if not cands:
        ws_c.cell(row=r, column=1, value=COUNTY)
        ws_c.cell(row=r, column=2, value=race)
        ws_c.cell(row=r, column=3, value="")
        ws_c.cell(row=r, column=4, value="")
        ws_c.cell(row=r, column=5, value="")
        ws_c.cell(row=r, column=6, value=CANDIDATES_SOURCE)
        r += 1
        empty_races += 1
        continue

    flag = flag_override
    if flag is None:
        uncontested = len(cands) <= seats
        flag = f"Uncontested - {len(cands)} candidate(s) for {seats} seat(s) per registrar" if uncontested else ""

    for name, occupation in cands:
        ws_c.cell(row=r, column=1, value=COUNTY)
        ws_c.cell(row=r, column=2, value=race)
        ws_c.cell(row=r, column=3, value=name)
        ws_c.cell(row=r, column=4, value=occupation)
        ws_c.cell(row=r, column=5, value=flag)
        ws_c.cell(row=r, column=6, value=CANDIDATES_SOURCE)
        r += 1
        cand_rows += 1

print(f"Wrote {cand_rows} candidate rows + {empty_races} empty-race rows across {len(RACES)} races")

wb.save(XLSX_PATH)
print("Saved.")
