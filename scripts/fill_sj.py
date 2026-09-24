# -*- coding: utf-8 -*-
import json
import openpyxl

XLSX_PATH = "/Users/ciarazavala/Desktop/Ballot Preview Backfill/Measure_Candidate Backfill Info.xlsx"
COUNTY = "San Joaquin County"
CANDIDATES_URL = ("https://www.sjgov.org/docs/default-source/registrar-of-voters-documents/elections/"
                   "gubernatorial-general-election-2026/candidate-services/candidate-roster-(35).pdf?sfvrsn=8b61f2ae_98")

UNCONTESTED_FLAG = ('Uncontested - registrar marks this contest "On Ballot: No" (filed candidates do not '
                    'exceed available seats, so no election is held for this office)')

CROSS_COUNTY = {
    '2519': 'Stanislaus County Board of Education, not San Joaquin County',
    '2550': 'Yosemite Community College District is based in Stanislaus County (Modesto)',
    '6050': 'Oakdale Irrigation District is based in Stanislaus County',
    '6065': 'Byron-Bethany Irrigation District is based in Contra Costa/Alameda County',
    '6066': 'Byron-Bethany Irrigation District is based in Contra Costa/Alameda County',
}

data = json.load(open('/private/tmp/claude-501/-Users-ciarazavala-Desktop-Ballot-Preview-Backfill/28cd04c7-f986-44cf-90af-3f056f743036/scratchpad/sj_races_parsed.json'))

wb = openpyxl.load_workbook(XLSX_PATH)
ws_c = wb['Candidates']

c_row = 2
while ws_c.cell(row=c_row, column=1).value:
    c_row += 1

n = 0
for r in data:
    flags = []
    if not r['on_ballot']:
        flags.append(UNCONTESTED_FLAG)
    if r['id'] in CROSS_COUNTY:
        flags.append(f'Cross-county jurisdiction: {CROSS_COUNTY[r["id"]]} - verify before publishing under "San Joaquin County"')
    flag = " ; ".join(flags)
    for c in r['candidates']:
        occ = c['occupation']
        if occ == 'No Ballot Designation':
            occ = ''
        ws_c.cell(row=c_row, column=1, value=COUNTY)
        ws_c.cell(row=c_row, column=2, value=r['race'])
        ws_c.cell(row=c_row, column=3, value=c['name'])
        ws_c.cell(row=c_row, column=4, value=occ)
        ws_c.cell(row=c_row, column=5, value=flag)
        ws_c.cell(row=c_row, column=6, value=CANDIDATES_URL)
        c_row += 1
        n += 1

wb.save(XLSX_PATH)
print(f"Candidates rows written: {n}")
