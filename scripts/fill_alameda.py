import json
import openpyxl

XLSX_PATH = "/Users/ciarazavala/Desktop/Ballot Preview Backfill/Measure_Candidate Backfill Info.xlsx"

COUNTY = "Alameda County"

# ---- Measures ----
measures = json.load(open('/private/tmp/claude-501/-Users-ciarazavala-Desktop-Ballot-Preview-Backfill/28cd04c7-f986-44cf-90af-3f056f743036/scratchpad/measures_parsed.json'))

# ---- Candidates ----
candidates = json.load(open('/private/tmp/claude-501/-Users-ciarazavala-Desktop-Ballot-Preview-Backfill/28cd04c7-f986-44cf-90af-3f056f743036/scratchpad/candidates_local.json'))

wb = openpyxl.load_workbook(XLSX_PATH)
ws_m = wb['Measures']
ws_c = wb['Candidates']

# find first empty row in Measures (row 2 onward)
m_row = 2
while ws_m.cell(row=m_row, column=1).value:
    m_row += 1

for m in measures:
    flag = ""
    if not m['jurisdiction']:
        flag = "No specific jurisdiction named by registrar beyond measure title (multi-county regional measure) - verify jurisdiction wording"
    ws_m.cell(row=m_row, column=1, value=COUNTY)
    ws_m.cell(row=m_row, column=2, value=m['title'])
    ws_m.cell(row=m_row, column=3, value=m['jurisdiction'] or "")
    ws_m.cell(row=m_row, column=4, value=m['desc'])
    ws_m.cell(row=m_row, column=5, value=flag)
    m_row += 1

# find first empty row in Candidates (row 2 onward)
c_row = 2
while ws_c.cell(row=c_row, column=1).value:
    c_row += 1

for r in candidates:
    race = r['race_title']
    if not r['candidates']:
        ws_c.cell(row=c_row, column=1, value=COUNTY)
        ws_c.cell(row=c_row, column=2, value=race)
        ws_c.cell(row=c_row, column=3, value="")
        ws_c.cell(row=c_row, column=4, value="")
        ws_c.cell(row=c_row, column=5, value="")
        c_row += 1
        continue
    for c in r['candidates']:
        ws_c.cell(row=c_row, column=1, value=COUNTY)
        ws_c.cell(row=c_row, column=2, value=race)
        ws_c.cell(row=c_row, column=3, value=c['name'])
        ws_c.cell(row=c_row, column=4, value=c['ballot_designation'])
        ws_c.cell(row=c_row, column=5, value="")
        c_row += 1

wb.save(XLSX_PATH)
print(f"Measures rows written: {len(measures)}; Candidates rows written: {sum(max(len(r['candidates']),1) for r in candidates)}")
