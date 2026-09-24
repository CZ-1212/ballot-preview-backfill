import json
import openpyxl

XLSX_PATH = "/Users/ciarazavala/Desktop/Ballot Preview Backfill/Measure_Candidate Backfill Info.xlsx"
COUNTY = "Contra Costa County"

CANDIDATE_SOURCE = ("https://www.contracostavote.gov/election/november-3-2026-general-election/#Candidates ; "
                    "https://www.contracostavote.gov/wp-content/uploads/8-27-26_candidatelist_detail_final.pdf")
MEASURE_SOURCE = "https://www.contracostavote.gov/wp-content/uploads/26Nov03_Measure-Wording-List-1.pdf"

candidates = json.load(open('/private/tmp/claude-501/-Users-ciarazavala-Desktop-Ballot-Preview-Backfill/28cd04c7-f986-44cf-90af-3f056f743036/scratchpad/cc_final_candidates.json'))
measures = json.load(open('/private/tmp/claude-501/-Users-ciarazavala-Desktop-Ballot-Preview-Backfill/28cd04c7-f986-44cf-90af-3f056f743036/scratchpad/cc_measures_parsed.json'))

wb = openpyxl.load_workbook(XLSX_PATH)
ws_m = wb['Measures']
ws_c = wb['Candidates']

m_row = 2
while ws_m.cell(row=m_row, column=1).value:
    m_row += 1

for m in measures:
    flag = "" if m['jurisdiction'] else "No specific jurisdiction named by registrar beyond measure title (multi-county regional measure) - verify jurisdiction wording"
    ws_m.cell(row=m_row, column=1, value=COUNTY)
    ws_m.cell(row=m_row, column=2, value=f"Measure {m['letter']}")
    ws_m.cell(row=m_row, column=3, value=m['jurisdiction'] or "")
    ws_m.cell(row=m_row, column=4, value=m['description'])
    ws_m.cell(row=m_row, column=5, value=flag)
    ws_m.cell(row=m_row, column=6, value=MEASURE_SOURCE)
    m_row += 1

c_row = 2
while ws_c.cell(row=c_row, column=1).value:
    c_row += 1

for c in candidates:
    ws_c.cell(row=c_row, column=1, value=COUNTY)
    ws_c.cell(row=c_row, column=2, value=c['office_name'])
    ws_c.cell(row=c_row, column=3, value=c['candidate'])
    ws_c.cell(row=c_row, column=4, value=c['occupation'])
    ws_c.cell(row=c_row, column=5, value=c.get('flag', ''))
    ws_c.cell(row=c_row, column=6, value=CANDIDATE_SOURCE)
    c_row += 1

wb.save(XLSX_PATH)
print(f"Measures rows written: {len(measures)}; Candidates rows written: {len(candidates)}")
