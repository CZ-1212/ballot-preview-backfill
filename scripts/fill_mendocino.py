# -*- coding: utf-8 -*-
import openpyxl

XLSX_PATH = "/Users/ciarazavala/Desktop/Ballot Preview Backfill/Measure_Candidate Backfill Info.xlsx"
COUNTY = "Mendocino County"

RACES_PDF = "https://www.mendocinocounty.gov/home/showpublisheddocument/79352/639229099880630000"
MEASURES_PAGE = "https://www.mendocinocounty.gov/government/assessor-county-clerk-recorder-elections/vitals-fees/measures-on-the-ballot"

OFFICE_TITLE_FLAG = ("Source (countywide candidate list) names the jurisdiction and seat count but not a "
                     "formal office title - \"Council Member\"/\"Director\"/\"Governing Board Member\" "
                     "inferred from standard CA local government structure, not printed verbatim - verify")

RACES = [
    ("Supervisor, 3rd District", "", [
        ("Buffey Wright Bourassa", "Program Manager"),
        ("Eric Hart", "Nonprofit/Business Advisor"),
    ]),
    ("City of Fort Bragg, Council Member", OFFICE_TITLE_FLAG, [
        ("Jason Godeke", "Incumbent"),
        ("Marcia Rafanan", "Incumbent"),
        ("Tess Albin-Smith", "Incumbent"),
        ("Mary Rose Kaczorowski", "Journalist/Artist"),
        ("Scott J. Mayberry", "Task Force Commander"),
        ("Levi W. Cooke", "College Student"),
        ("Mary Ann Schmidt", "Retired Teacher"),
        ("Richard Uleving", "Business Owner"),
    ]),
    ("City of Point Arena, Council Member", OFFICE_TITLE_FLAG, [
        ("Barbara Burkey", "Incumbent"),
        ("Joshua Karlin Resnick", "City Planner"),
    ]),
    ("City of Ukiah, Council Member", OFFICE_TITLE_FLAG, [
        ("Susan Sher", "Ukiah City Councilmember"),
        ("Juan V. Orozco", "Incumbent"),
        ("Mari Rodin", "Incumbent"),
        ("Sean K. White", "Water Resources Manager"),
        ("Robert Bazzani", "Small Business Owner"),
        ("Mia L. Uribe", "Retail Associate"),
        ("John Strangio", "Business Owner"),
        ("Matt Froneberger", "Chief Plant Operator"),
    ]),
    ("City of Willits, Council Member", OFFICE_TITLE_FLAG, [
        ("Robin Leler", "Teacher"),
        ("Sherrie Ebyam", "Retired Finance Director"),
    ]),
    ("Potter Valley Irrigation District, Director, Division 4 - Short Term", OFFICE_TITLE_FLAG, [
        ("Steven Giuntini", "Appointed Incumbent"),
        ("Jarred Hildebrand", "Rancher/Vineyard Manager"),
    ]),
    ("Mendocino County Board of Education, Governing Board Member, Trustee Area 1", OFFICE_TITLE_FLAG, [
        ("Marilyn Puget", "Incumbent"),
        ("Craig Perry", "Accountant"),
    ]),
    ("Ukiah Unified School District, Governing Board Member, Trustee Area 1", OFFICE_TITLE_FLAG, [
        ('Beatriz "Bea" Arkin', "Incumbent"),
        ("Gabriel Baca Meza", "Caregiver"),
    ]),
    ("Willits Unified School District, Governing Board Member", OFFICE_TITLE_FLAG, [
        ("Jeanne King", "Incumbent"),
        ("Paula Nunez", "Incumbent"),
        ("Mason Neil Rescina", "Instructional Aide"),
        ("Nicole Karkar", "Business Owner"),
    ]),
]

# (letter, jurisdiction, description, source_url, flag)
MEASURES = [
    ("B", "City of Fort Bragg",
     "Fort Bragg Fire Protection Measure. To maintain reliable local fire protection services; replace aging fire engines and emergency response equipment; keep funds locally controlled and unavailable for State use; and require annual audit and public reporting, shall the City of Fort Bragg authorize a $50 annual parcel tax for 10 years, generating approximately $123,900 annually?",
     RACES_PDF + " ; " + MEASURES_PAGE + " ; https://www.mendocinocounty.gov/home/showpublisheddocument/79582/639239569418670000",
     ""),
    ("C", "City of Point Arena",
     "Shall Ordinance No. 247 amending Chapter 3.30.030 of the Point Arena Municipal Code to increase the City of Point Arena's existing Transient Occupancy Tax (a bed tax paid when overnight visitors rent a room) from twelve percent (12%) to fourteen percent (14%) to fund public services and maintain public areas, effective April 1, 2027, which proposed rate increase and amendment is anticipated to raise an additional $35,833 per year in revenue and which will continue until repealed by the City Council or the city voters, be adopted?",
     RACES_PDF + " ; " + MEASURES_PAGE + " ; https://www.mendocinocounty.gov/home/showpublisheddocument/79584/639239569421470000",
     ""),
    ("D", "Coast Life Support District",
     "Allows Coast Life Support District to increase the current special parcel tax rate on real property to $101 per unit of benefit, for emergency medical services, including ambulance, life support and medical transport purposes and other necessary operating expenses, providing an estimated $1,972,934 annually, until the tax is amended or repealed; authorizes annual increases for inflation; and raises the District's appropriations limit to permit spending of tax the revenue.",
     RACES_PDF + " ; " + MEASURES_PAGE + " ; https://www.mendocinocounty.gov/home/showpublisheddocument/79586/639239569423130000",
     'Cross-county jurisdiction: Coast Life Support District’s own ordinance text states the tax is collected by "the Counties of Mendocino and Sonoma" - district spans both counties - verify before publishing under "Mendocino County"'),
    ("AB", "Sonoma County Junior College District",
     "Without increasing tax rates, to improve facilities to train students for careers in nursing/firefighting/first responders, building trades, agriculture, auto mechanics; attract/retain high quality faculty, prepare students for universities; repair leaky roofs; repair, construct, and acquire classrooms/sites/facilities/equipment, shall the Sonoma County Junior College District measure be adopted authorizing $830,000,000 in bonds at legal rates, levying $25 per $100,000 of assessed valuation ($55,300,000 annually), while bonds are outstanding, and with oversight/audits/local control?",
     RACES_PDF + " ; " + MEASURES_PAGE + " ; https://www.mendocinocounty.gov/home/showpublisheddocument/79580/639239569415530000",
     'Cross-county jurisdiction: Sonoma County Junior College District (Santa Rosa Junior College), not Mendocino County - same measure also appears on Marin County’s ballot as Measure AB - verify before publishing under "Mendocino County"'),
    ("K", "Southern Humboldt Community Healthcare District",
     "To replace 77 year-old Jerold Phelps Community Hospital with a modern hospital to provide life-saving emergency care for victims of accidents, heart attacks, strokes, and other emergencies, improve Emergency Room facilities, operating rooms, provide advanced medical equipment/technology and life flight services; and qualify for matching funds shall Southern Humboldt Healthcare District's measure authorizing $25,000,000 in bonds be adopted, levying 8 cents per $100 assessed value (raising $2,000,000 annually) while bonds are outstanding, with independent citizen oversight and all funds staying local?",
     RACES_PDF + " ; " + MEASURES_PAGE + " ; https://www.mendocinocounty.gov/home/showpublisheddocument/79588/639239569425470000",
     'Cross-county jurisdiction: Southern Humboldt Community Healthcare District is primarily in Humboldt County (district’s own filing describes it as serving "Southern Humboldt County and Mendocino County") - verify before publishing under "Mendocino County"'),
]

wb = openpyxl.load_workbook(XLSX_PATH)
ws_m = wb['Measures']
ws_c = wb['Candidates']

m_row = 2
while ws_m.cell(row=m_row, column=1).value:
    m_row += 1

n_measures = 0
for letter, jurisdiction, desc, url, flag in MEASURES:
    ws_m.cell(row=m_row, column=1, value=COUNTY)
    ws_m.cell(row=m_row, column=2, value=f"Measure {letter}")
    ws_m.cell(row=m_row, column=3, value=jurisdiction)
    ws_m.cell(row=m_row, column=4, value=desc)
    ws_m.cell(row=m_row, column=5, value=flag)
    ws_m.cell(row=m_row, column=6, value=url)
    m_row += 1
    n_measures += 1

c_row = 2
while ws_c.cell(row=c_row, column=1).value:
    c_row += 1

n_candidates = 0
for race, race_flag, cands in RACES:
    for name, occ in cands:
        ws_c.cell(row=c_row, column=1, value=COUNTY)
        ws_c.cell(row=c_row, column=2, value=race)
        ws_c.cell(row=c_row, column=3, value=name)
        ws_c.cell(row=c_row, column=4, value=occ)
        ws_c.cell(row=c_row, column=5, value=race_flag)
        ws_c.cell(row=c_row, column=6, value=RACES_PDF)
        c_row += 1
        n_candidates += 1

wb.save(XLSX_PATH)
print(f"Measures rows written: {n_measures}; Candidates rows written: {n_candidates}")
