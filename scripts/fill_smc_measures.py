# -*- coding: utf-8 -*-
import openpyxl

XLSX_PATH = "/Users/ciarazavala/Desktop/Ballot Preview Backfill/Measure_Candidate Backfill Info.xlsx"
COUNTY = "San Mateo County"
MEASURES_PAGE = "https://smcacre.gov/elections/november-3-2026-statewide-general-election"

CORRUPTED_FLAG = ('Could not confirm exact ballot question wording: this measure’s source PDF has a '
                   'corrupted/unreliable text layer (common words like "shall", "YES", "NO" do not extract '
                   'cleanly - likely a copy-protected or scanned document) - left blank rather than guess; '
                   'would need manual visual review of the source PDF to transcribe accurately')

# (letter, jurisdiction, description, source_url, flag)
FOUND = [
    ("L", "County of San Mateo",
     "Shall the San Mateo County Charter be amended to add a provision to the Charter Preamble designating long-term resilience to extreme weather events as a County priority?",
     "https://smcacre.gov/system/files/2026-07/San%20Mateo%20County%20Amendments.pdf", ""),
    ("AA", "County of San Mateo",
     "Shall the San Mateo County Charter be amended to add a provision to the Charter Preamble affirming the County’s commitment to the dignity, civil rights, and equal protection of all County residents?",
     "https://smcacre.gov/system/files/2026-07/San%20Mateo%20County%20Amendments.pdf", ""),
    ("J", "County of San Mateo",
     "Shall the San Mateo County Charter be amended to extend the period in which the County Board of Supervisors is required to either call a special election or make an appointment in the event of a vacancy in an elected office other than the Board of Supervisors from 30 days to 60 days, and to require the Board to conduct at least one public meeting to receive public comment prior to any appointment?",
     "https://smcacre.gov/system/files/2026-07/San%20Mateo%20County%20Amendments.pdf", ""),
    ("U", "County of San Mateo",
     "Shall the San Mateo County Charter be amended to require that supervisorial districts be reapportioned following each federal decennial census by an independent redistricting commission pursuant to applicable provisions of the California Elections Code?",
     "https://smcacre.gov/system/files/2026-07/San%20Mateo%20County%20Amendments.pdf", ""),
    ("RTM", "",
     "To prevent major service cuts to BART and other transit, avoid increased traffic, and reduce pollution by: Preserving BART, Caltrain, VTA, SamTrans, AC Transit, Muni, other transit for everyone, including workers, students, seniors, persons with disabilities; Supporting transit safety, cleanliness, affordability, reliability; Repairing targeted roads/potholes; Requiring financial transparency, oversight, accountability; shall the measure enacting a 0.5% (Alameda, Contra Costa, San Mateo, Santa Clara counties), and 1% (San Francisco) sales tax for 14 years generating approximately $980,000,000 annually, be adopted?",
     "https://smcacre.gov/system/files/2026-07/52-ENG-M206-Contest%20Code-Reso-Public%20Transit%20District%20Sales%20Tax_Redacted.pdf",
     'Regional measure spanning 5 Bay Area counties (Alameda, Contra Costa, San Francisco, San Mateo, Santa Clara) - no single-jurisdiction name given by the registrar beyond the title'),
    ("K", "San Carlos School District",
     "To protect outstanding educational programs, hands-on science, math, technology, arts, reading and writing, attract and retain high-quality teachers, maintain reasonable class sizes, and sustain property values, shall the San Carlos School District’s measure be adopted to replace its existing parcel tax at a $489 per parcel rate (raising approximately $4,400,000 annually) for 8 years, with annual cost-of-living adjustments, citizen oversight, senior exemptions, all funds spent locally and no funds used for administrator salaries?",
     "https://smcacre.gov/system/files/2026-05/52-ENG-M101-Contest%20Code-Reso-San%20Carlos%20SD%20Parcel%20Tax_Redacted.pdf", ""),
    ("M", "Jefferson Union High School District",
     "With funds that cannot be taken by the State or Federal Government and spent elsewhere, shall Jefferson Union High School District’s measure to directly support teachers and academic programs; expand / enhance career training and college preparedness classes; and hire and retain qualified teachers and counselors at Jefferson High School, Oceana High School, Shasta High School, Terra Nova High School, Thornton High School, and Westmoor High School be adopted, replacing the 2018 parcel tax at $98 per parcel for ten years (raising $3.4 million annually) with full public disclosure of all spending?",
     "https://smcacre.gov/system/files/2026-07/52-ENG-M106-Contest%20Code-Reso-JUHSD%20Parcel%20Tax_Redacted.pdf", ""),
    ("CC", "City of East Palo Alto",
     "Shall the measure providing that no person shall be eligible to serve for more than three (3) consecutive full terms on the City Council be adopted, with service of partial terms exceeding two years qualifying as a full term?",
     "https://smcacre.gov/system/files/2026-07/52-ENG-M207-Contest%20-Code-Reso-EPA%20Term%20Limits_Redacted.pdf", ""),
    ("O", "City of East Palo Alto",
     "East Palo Alto Community/Water/Public Safety Infrastructure Measure. Shall the measure constructing a Library, Emergency Command Center, civic center, youth sports park; replacing deteriorating water pipes to provide safe/clean drinking water; repairing storm drains to prevent flooding; addressing earthquake safety, by authorizing East Palo Alto to issue $125,000,000 in general obligation bonds, with annual levies averaging $105.26 per $100,000 of assessed value (raising approximately $8,940,000 annually) while bonds are outstanding, with independent audits/public disclosure, be adopted?",
     "https://smcacre.gov/system/files/2026-07/52-ENG-M207-Contest%20Code-Reso-City%20of%20East%20Palo%20Alto%20Bond_Redacted.pdf", ""),
    ("G", "City of San Carlos",
     "To provide funding for San Carlos city services, such as keeping local streets, sidewalks, storm drains, parks, and community facilities safe, clean, and well-maintained; fixing streets and potholes; relieving traffic congestion; repairing aging infrastructure; and supporting fire protection, paramedic, crime prevention, and 911 emergency response; shall City of San Carlos’ ordinance establishing a one-half cent sales tax be adopted, providing $6,000,000 annually for general government use for 14 years, with independent audits and all funds locally-controlled?",
     "https://smcacre.gov/system/files/2026-07/52-ENG-M104-Contest%20Code-Reso-City%20of%20San%20Carlos%20Use%20Tax_Redacted.pdf", ""),
    ("P", "City of Menlo Park",
     "Should the measure, which prohibits the City of Menlo Park from making changes to City-owned downtown parking lot properties, including selling, leasing, or conveying any of the properties, authorizing a new use such as housing development on any of the properties, or modifying, altering, or constructing improvements on any of the properties, if such change would diminish parking availability, access or convenience, unless the voters pass another ballot measure to approve the change, be adopted?",
     "https://smcacre.gov/system/files/2026-05/52-ENG-M102-Contest%20Code-Reso-Menlo%20Park%20Muni%20Amendment_Redacted.pdf", ""),
    ("E", "City of Redwood City",
     "Shall the measure repealing Redwood City Municipal Code Chapter 42 and replacing it with a new Chapter 42 that (1) limits rent increases on multifamily housing built before February 1, 1995; (2) expands just-cause eviction protections to most rental housing, including single-family homes, ADUs, and affordable housing; (3) requires new City administration estimated to cost $5 million to $11 million annually, funded by annual rental property owner fees; and (4) increases relocation assistance for no-fault evictions, be adopted?",
     "https://smcacre.gov/system/files/2026-07/52-ENG-M108-Contest%20Code-Reso-RWC%20Housing_Redacted.pdf", ""),
    ("X", "City of Brisbane",
     "Shall the measure updating the City's Business License Tax by replacing the current tiered structure with a $50 tax for businesses with annual gross receipts of $25,000 or less and $50 plus $1.00 per $1,000 of gross receipts above $25,000 for most other businesses; establishing a minimum tax base equal to the cost of operations for businesses whose operating costs exceed gross receipts; and raising approximately $1,000,000 annually, until repealed by voters, be adopted?",
     "https://smcacre.gov/system/files/2026-07/52-ENG-M302-Contest%20Code-Reso-City%20of%20Brisbane%20Business%20License%20Tax_Redacted_1.pdf", ""),
]

# (letter, jurisdiction) for the 16 blocked by corrupted text layer
BLOCKED = [
    ("R", "Burlingame School District"),
    ("Z", "Cabrillo Unified School District"),
    ("T", "Pacifica School District"),
    ("D", "City of Burlingame"),
    ("N", "Town of Portola Valley"),
    ("F", "City of San Mateo"),
    ("H", "City of San Mateo"),
    ("S", "City of San Bruno"),
    ("DD", "City of San Bruno"),
    ("EE", "City of San Bruno"),
    ("W", "San Mateo-Foster City School District"),
    ("V", "San Mateo County Community College District"),
    ("Y", "City of Belmont"),
    ("I", "City of Half Moon Bay"),
    ("Q", "City of Half Moon Bay"),
    ("BB", "City of Pacifica"),
]

wb = openpyxl.load_workbook(XLSX_PATH)
ws_m = wb['Measures']

m_row = 2
while ws_m.cell(row=m_row, column=1).value:
    m_row += 1

n = 0
for letter, jurisdiction, desc, url, flag in FOUND:
    ws_m.cell(row=m_row, column=1, value=COUNTY)
    ws_m.cell(row=m_row, column=2, value=f"Measure {letter}")
    ws_m.cell(row=m_row, column=3, value=jurisdiction)
    ws_m.cell(row=m_row, column=4, value=desc)
    ws_m.cell(row=m_row, column=5, value=flag)
    ws_m.cell(row=m_row, column=6, value=url + " ; " + MEASURES_PAGE)
    m_row += 1
    n += 1

for letter, jurisdiction in BLOCKED:
    ws_m.cell(row=m_row, column=1, value=COUNTY)
    ws_m.cell(row=m_row, column=2, value=f"Measure {letter}")
    ws_m.cell(row=m_row, column=3, value=jurisdiction)
    ws_m.cell(row=m_row, column=4, value="")
    ws_m.cell(row=m_row, column=5, value=CORRUPTED_FLAG)
    ws_m.cell(row=m_row, column=6, value=MEASURES_PAGE)
    m_row += 1
    n += 1

wb.save(XLSX_PATH)
print(f"Measures rows written: {n} ({len(FOUND)} with confirmed text, {len(BLOCKED)} blocked/flagged)")
