import sys
sys.path.insert(0, '/private/tmp/claude-501/-Users-ciarazavala-Desktop-Ballot-Preview-Backfill/28cd04c7-f986-44cf-90af-3f056f743036/scratchpad')
from sf_data import RACES, MEASURES, CANDIDATES_URL, MEASURES_URL
import openpyxl

XLSX_PATH = "/Users/ciarazavala/Desktop/Ballot Preview Backfill/Measure_Candidate Backfill Info.xlsx"
COUNTY = "City and County of San Francisco"

wb = openpyxl.load_workbook(XLSX_PATH)
ws_m = wb['Measures']
ws_c = wb['Candidates']

m_row = 2
while ws_m.cell(row=m_row, column=1).value:
    m_row += 1

n_measures = 0
for letter, title, jurisdiction, desc, flag in MEASURES:
    ws_m.cell(row=m_row, column=1, value=COUNTY)
    ws_m.cell(row=m_row, column=2, value=f"Measure {letter} - {title}")
    ws_m.cell(row=m_row, column=3, value=jurisdiction)
    ws_m.cell(row=m_row, column=4, value=desc)
    ws_m.cell(row=m_row, column=5, value=flag)
    ws_m.cell(row=m_row, column=6, value=MEASURES_URL)
    m_row += 1
    n_measures += 1

c_row = 2
while ws_c.cell(row=c_row, column=1).value:
    c_row += 1

n_candidates = 0
for race, cands in RACES:
    for name, occ in cands:
        ws_c.cell(row=c_row, column=1, value=COUNTY)
        ws_c.cell(row=c_row, column=2, value=race)
        ws_c.cell(row=c_row, column=3, value=name)
        ws_c.cell(row=c_row, column=4, value=occ)
        ws_c.cell(row=c_row, column=5, value="")
        ws_c.cell(row=c_row, column=6, value=CANDIDATES_URL)
        c_row += 1
        n_candidates += 1

wb.save(XLSX_PATH)
print(f"Measures rows written: {n_measures}; Candidates rows written: {n_candidates}")
