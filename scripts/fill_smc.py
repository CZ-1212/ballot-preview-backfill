# -*- coding: utf-8 -*-
import json
import openpyxl

XLSX_PATH = "/Users/ciarazavala/Desktop/Ballot Preview Backfill/Measure_Candidate Backfill Info.xlsx"
COUNTY = "San Mateo County"
CANDIDATES_URL = ("https://smcacre.gov/archival-document?document=https%3A%2F%2Fsmcacre.gov%2Fsystem%2Ffiles%2F"
                   "2026-09%2F52_candidateroster0903.pdf&title=Roster+of+Candidates")

UNCONTESTED_FLAG = ('Uncontested - registrar marks this contest "On Ballot: No" (filed candidates do not '
                    'exceed available seats, so no election is held for this office)')

CROSS_COUNTY = {
    '6260': 'Midpeninsula Regional Open Space District is based in Los Altos, Santa Clara County',
    '6261': 'Midpeninsula Regional Open Space District is based in Los Altos, Santa Clara County',
}

data = json.load(open('/private/tmp/claude-501/-Users-ciarazavala-Desktop-Ballot-Preview-Backfill/28cd04c7-f986-44cf-90af-3f056f743036/scratchpad/smc_local_races.json'))

wb = openpyxl.load_workbook(XLSX_PATH)
ws_c = wb['Candidates']

c_row = 2
while ws_c.cell(row=c_row, column=1).value:
    c_row += 1

n = 0
for r in data:
    flags = []
    if not r['on_ballot']:
        flags.append(UNCONTESTED_FLAG)
    if r['id'] in CROSS_COUNTY:
        flags.append(f'Cross-county jurisdiction: {CROSS_COUNTY[r["id"]]} - verify before publishing under "San Mateo County"')
    flag = " ; ".join(flags)
    for c in r['candidates']:
        occ = c['occupation']
        if occ == '(No Ballot Designation)':
            occ = ''
        ws_c.cell(row=c_row, column=1, value=COUNTY)
        ws_c.cell(row=c_row, column=2, value=r['race'])
        ws_c.cell(row=c_row, column=3, value=c['name'])
        ws_c.cell(row=c_row, column=4, value=occ)
        ws_c.cell(row=c_row, column=5, value=flag)
        ws_c.cell(row=c_row, column=6, value=CANDIDATES_URL)
        c_row += 1
        n += 1

wb.save(XLSX_PATH)
print(f"Candidates rows written: {n}")
